import openpyxl

# --- CONFIG ---
EXCEL_PATH = "Ms4300_Excel___Copie.xlsx"  # cambia col nome del file caricato su Colab
SHEET_NAME = "noms"

# Ripiego SOLO per i coniugi non catalogati che non hanno nemmeno un link
# ipertestuale nella cella lien_conjoint_X. Aggiungi qui se ne trovi altri a mano.
EXTERNAL_URIS_FALLBACK = {
    # "Chézeaux, Catherine de": "https://...",
}

# --- CARICAMENTO DATI ---
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb[SHEET_NAME]
headers = [c.value for c in ws[1]]

def col(name):
    return headers.index(name) + 1

id_c = col('id')
name_c = col('nom_standard')

# nome -> id, per risolvere il nome di un coniuge al suo id ITIE quando ha una riga propria
name_to_id = {}
for r in range(2, ws.max_row + 1):
    nm = ws.cell(row=r, column=name_c).value
    idv = ws.cell(row=r, column=id_c).value
    if nm not in (None, 'null'):
        name_to_id[nm.strip()] = idv

# --- RACCOLTA MATRIMONI ---
marriages = []
seen_pairs = set()  # coppie {id_a, id_b} già trasformate in un evento, per non duplicare

for slot_conj, slot_lien, slot_date in [
    ('conjoint_1', 'lien_conjoint_1', 'date_mariage_1'),
    ('conjoint_2', 'lien_conjoint_2', 'date_mariage_2'),
]:
    conj_c = col(slot_conj)
    lien_c = col(slot_lien)
    date_c = col(slot_date)
    for r in range(2, ws.max_row + 1):
        my_id = ws.cell(row=r, column=id_c).value
        spouse_name = ws.cell(row=r, column=conj_c).value
        date_val = ws.cell(row=r, column=date_c).value
        if spouse_name in (None, 'null'):
            continue
        spouse_name = spouse_name.strip()
        date_val = None if date_val in (None, 'null') else str(date_val)

        spouse_id = name_to_id.get(spouse_name)

        if spouse_id:
            pair = frozenset({my_id, spouse_id})
            if pair in seen_pairs:
                continue
            seen_pairs.add(pair)
            marriages.append({
                'partner_a': my_id,
                'partner_b': spouse_id,
                'external': None,
                'name_hint': None,
                'date': date_val,
            })
        else:
            # non catalogato: guarda se la cella lien_conjoint_X ha un link nascosto
            lien_cell = ws.cell(row=r, column=lien_c)
            external_uri = lien_cell.hyperlink.target if lien_cell.hyperlink else EXTERNAL_URIS_FALLBACK.get(spouse_name)
            marriages.append({
                'partner_a': my_id,
                'partner_b': None,
                'external': external_uri,
                'name_hint': spouse_name,
                'date': date_val,
            })

print(f"Totale matrimoni trovati: {len(marriages)}")
con_uri_esterno = sum(1 for m in marriages if m['external'])
con_blank_node = sum(1 for m in marriages if not m['partner_b'] and not m['external'])
print(f"  di cui con URI esterno recuperato: {con_uri_esterno}")
print(f"  di cui ancora come blank node (nessun link trovato): {con_blank_node}")

# --- GENERAZIONE TURTLE ---
lines = []
width = max(3, len(str(len(marriages))))
for i, m in enumerate(marriages, start=1):
    event_id = f"ITIE_M{i:0{width}d}"
    lines.append(f"marriage:{event_id} rdf:type bio:Marriage ;")
    lines.append(f"    bio:partner person:{m['partner_a']} ;")
    if m['partner_b']:
        lines.append(f"    bio:partner person:{m['partner_b']} ;")
    elif m['external']:
        lines.append(f"    bio:partner <{m['external']}> ;")
    else:
        lines.append(f"    bio:partner [ rdf:type foaf:Person ; schema:name \"{m['name_hint']}\" ] ;")
    if m['date']:
        lines.append(f'    bio:date "{m["date"]}" .')
    else:
        lines[-1] = lines[-1].rstrip(' ;') + ' .'
    lines.append("")

ttl_output = "\n".join(lines)
with open("matrimoni.ttl", "w", encoding="utf-8") as f:
    f.write(ttl_output)

print("File 'matrimoni.ttl' scritto.")
